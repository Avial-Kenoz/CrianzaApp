"""
Importa consumo mensual de alimento desde planillas Excel historicas (2023-2026).
Lee la hoja "Total Consumo Mensual" de cada archivo y carga:
  - feed_monthly_consumption_legacy : kg por estanque y tipo de alimento
  - feed_monthly_stock_legacy        : stock inicial, ingresos, consumido, final

Uso:
    python scripts/import_alimentacion_legacy.py            # carga todo
    python scripts/import_alimentacion_legacy.py --dry-run  # solo imprime, no inserta
    python scripts/import_alimentacion_legacy.py --year 2024
"""

import os
import sys
import argparse
import re
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")

from app.models.feed import FeedMonthlyConsumptionLegacy, FeedMonthlyStockLegacy
from app.db.session import Base

DATABASE_URL = os.environ["DATABASE_URL"]

# ---------------------------------------------------------------------------
# Normalización de nombres de tipos de alimento
# Cada entrada: patron (substring, case-insensitive) -> feed_type_key
# Se evalúan en orden; se usa el PRIMER match.
# ---------------------------------------------------------------------------
FEED_TYPE_PATTERNS = [
    # Infa pequeños
    ("infa 0.2/0.4",   "infa_combinado"),   # columna doble — se trata como combinado
    ("0,2 mm infa",    "infa_0.2mm"),
    ("0.2 infa",       "infa_0.2mm"),
    ("0,4 mm infa",    "infa_0.4mm"),
    ("0.4 infa",       "infa_0.4mm"),
    # Thalassa
    ("thalassa 0.5",   "thalassa_0.5mm"),
    ("0,5 mm thalassa","thalassa_0.5mm"),
    ("thalassa 0.9",   "thalassa_0.9mm"),
    ("thalassa 1.3",   "thalassa_1.3mm"),
    ("1,3 mm thalassa","thalassa_1.3mm"),
    ("1.3 thalassa",   "thalassa_1.3mm"),
    # Progress / Metabolica
    ("3 mm progress",  "progress_3mm"),
    ("4,5 mm metab",   "progress_4.5mm"),
    ("4.5 mm progress","progress_4.5mm"),
    ("4.5 progress",   "progress_4.5mm"),
    ("6.0 mm progress","progress_6mm"),
    ("6 mm metabolica","progress_6mm"),
    ("6 mm metab",     "progress_6mm"),
    ("6 mm progress",  "progress_6mm"),
    ("8 mm metab",     "progress_8mm"),
    ("8 mm progress",  "progress_8mm"),
    # Sturgeon / Reproductor
    ("6 mm sturgeon",  "sturgeon_rep_6mm"),
    ("8 mm sturgeon",  "reproductor_8mm"),
    ("8 mm reproductor","reproductor_8mm"),
    # Ewos
    ("ewos micro psz le 100", "ewos_psz_4mm"),
    ("ewos micro psz le 250", "ewos_psz_6mm"),
    ("ewos breed",     "ewos_breed_9.5mm"),
]

MONTH_NAMES = {
    "enero":1,"febrero":2,"marzo":3,"abril":4,"mayo":5,"junio":6,
    "julio":7,"agosto":8,"septiembre":9,"octubre":10,"noviembre":11,"diciembre":12,
}

POND_CODE_RE = re.compile(r"^(SO|SP|NO|NC|NP|C)\d+[A-Z]?$", re.IGNORECASE)

SKIP_ROW_KEYWORDS = ("hatchery","incubaci","alevinaje","jaula","ext exterior")
TOTAL_ROW_KEYWORDS = ("total consumido","inicial mensual","ingresos","final mensual")


def normalize_feed_type(raw: str) -> str | None:
    low = raw.lower().strip()
    for pattern, key in FEED_TYPE_PATTERNS:
        if pattern.lower() in low:
            return key
    return None


def extract_month_year(ws, filename: str, year_from_dir: int) -> tuple[int, int]:
    """
    Extrae (year, month). El año siempre viene del directorio.
    El mes se lee PRIMERO del nombre del archivo (más confiable — los headers a veces
    quedan con el mes anterior cuando se copia la plantilla sin actualizar).
    Solo si el filename no contiene mes se consulta la hoja.
    """
    # Primero: nombre del archivo
    fname_low = filename.lower()
    for name, num in MONTH_NAMES.items():
        if name in fname_low:
            return year_from_dir, num

    # Fallback: texto de mes en las primeras 8 filas de la hoja
    for row in ws.iter_rows(min_row=1, max_row=8, max_col=15, values_only=True):
        for cell in row:
            if cell and isinstance(cell, str) and cell.lower().strip() in MONTH_NAMES:
                return year_from_dir, MONTH_NAMES[cell.lower().strip()]

    raise ValueError(f"No se pudo determinar el mes en {filename}")


def find_header_row(ws) -> tuple[int, dict[int, tuple[str, str]]]:
    """
    Devuelve (header_row_index_1based, {col_index: (raw_name, feed_type_key)}).
    Busca la fila que tenga >= 2 tipos de alimento reconocibles.
    """
    for row_idx, row in enumerate(ws.iter_rows(min_row=8, max_row=18, values_only=True), start=8):
        feed_cols = {}
        for col_idx, cell in enumerate(row, start=1):
            if not cell or not isinstance(cell, str):
                continue
            key = normalize_feed_type(cell)
            if key:
                feed_cols[col_idx] = (cell.strip(), key)
        if len(feed_cols) >= 2:
            return row_idx, feed_cols
    raise ValueError("No se encontró fila de encabezado con tipos de alimento")


def is_pond_row(row: tuple) -> str | None:
    """Retorna el pond_code si la fila corresponde a un estanque, None si no."""
    if len(row) < 2:
        return None
    pond_code = row[1]
    if not pond_code or not isinstance(pond_code, str):
        return None
    pond_code = pond_code.strip()
    if POND_CODE_RE.match(pond_code):
        return pond_code
    return None


def is_total_row(row: tuple) -> str | None:
    """Retorna la clave de la fila de totales o None."""
    col_a = str(row[0]).lower().strip() if row[0] else ""
    for kw in TOTAL_ROW_KEYWORDS:
        if kw in col_a:
            return col_a
    return None


def should_skip_row(row: tuple) -> bool:
    col_a = str(row[0]).lower() if row[0] else ""
    col_b = str(row[1]).lower() if len(row) > 1 and row[1] else ""
    for kw in SKIP_ROW_KEYWORDS:
        if kw in col_a or kw in col_b:
            return True
    return False


def to_float(val) -> float | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val).replace(",", ".").replace(" ", ""))
    except (ValueError, TypeError):
        return None


def process_file(path: Path, dry_run: bool, session: Session) -> dict:
    year_from_dir = int(path.parent.name.split()[-1])
    stats = {"inserted_consumption": 0, "skipped_consumption": 0,
             "inserted_stock": 0, "skipped_stock": 0, "warnings": []}

    wb = openpyxl.load_workbook(path, data_only=True)

    # Buscar hoja de consumo mensual
    sheet = None
    for name in wb.sheetnames:
        if "consumo" in name.lower() and "mensual" in name.lower():
            sheet = wb[name]
            break
    if sheet is None:
        stats["warnings"].append("Sin hoja 'Total Consumo Mensual'")
        return stats

    try:
        year, month = extract_month_year(sheet, path.name, year_from_dir)
        header_row_idx, feed_cols = find_header_row(sheet)
    except ValueError as e:
        stats["warnings"].append(str(e))
        return stats

    source = str(path.relative_to(ROOT))

    # ---- Lectura de filas de datos (estanques) ----
    reached_totals = False

    for row_idx, row in enumerate(sheet.iter_rows(min_row=header_row_idx + 1,
                                                   max_row=sheet.max_row,
                                                   values_only=True), start=header_row_idx + 1):
        if not any(v is not None for v in row):
            continue

        total_key = is_total_row(row)
        if total_key:
            reached_totals = True

        if reached_totals:
            # Procesar filas de resumen de stock
            if total_key in ("total consumido", "inicial mensual", "ingresos", "final mensual"):
                for col_idx, (raw_name, feed_key) in feed_cols.items():
                    val = to_float(row[col_idx - 1]) if col_idx - 1 < len(row) else None
                    if val is None:
                        continue
                    # Buscar o crear registro de stock para este mes + tipo
                    existing = session.query(FeedMonthlyStockLegacy).filter_by(
                        year=year, month=month, feed_type_key=feed_key
                    ).first() if not dry_run else None

                    field_map = {
                        "total consumido": "consumed_kg",
                        "inicial mensual": "initial_kg",
                        "ingresos": "receipts_kg",
                        "final mensual": "final_kg",
                    }
                    field = field_map.get(total_key)
                    if not field:
                        continue

                    if dry_run:
                        stats["inserted_stock"] += 1
                        continue

                    if existing is None:
                        existing = FeedMonthlyStockLegacy(
                            year=year, month=month,
                            feed_type_key=feed_key, feed_type_raw=raw_name,
                            source_file=source,
                        )
                        session.add(existing)
                    setattr(existing, field, val)
                    stats["inserted_stock"] += 1
            continue

        if should_skip_row(row):
            continue

        pond_code = is_pond_row(row)
        if not pond_code:
            continue

        for col_idx, (raw_name, feed_key) in feed_cols.items():
            val = to_float(row[col_idx - 1]) if col_idx - 1 < len(row) else None
            if val is None or val == 0.0:
                continue

            if dry_run:
                print(f"  [{year}-{month:02d}] {pond_code:5s} {feed_key:<20s} {val:>8.1f} kg")
                stats["inserted_consumption"] += 1
                continue

            exists = session.query(FeedMonthlyConsumptionLegacy).filter_by(
                year=year, month=month, pond_code=pond_code, feed_type_key=feed_key
            ).first()
            if exists:
                stats["skipped_consumption"] += 1
                continue

            session.add(FeedMonthlyConsumptionLegacy(
                year=year, month=month,
                pond_code=pond_code,
                feed_type_raw=raw_name,
                feed_type_key=feed_key,
                consumed_kg=val,
                source_file=source,
            ))
            stats["inserted_consumption"] += 1

    if not dry_run:
        session.commit()

    return stats


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--year", type=int, default=None)
    args = parser.parse_args()

    engine = create_engine(DATABASE_URL)
    session = Session(engine)

    year_dirs = sorted(ROOT.glob("Alimentacion 202*/"))
    if args.year:
        year_dirs = [d for d in year_dirs if str(args.year) in d.name]

    total_consumed = 0
    total_stock = 0
    total_warnings = 0

    for year_dir in year_dirs:
        xlsx_files = sorted(year_dir.glob("*.xlsx"))
        for path in xlsx_files:
            label = f"{path.parent.name}/{path.name}"
            if args.dry_run:
                print(f"\n=== {label} ===")
            stats = process_file(path, args.dry_run, session)
            total_consumed += stats["inserted_consumption"]
            total_stock += stats["inserted_stock"]
            total_warnings += len(stats["warnings"])
            status = "OK" if not stats["warnings"] else f"WARN: {'; '.join(stats['warnings'])}"
            if not args.dry_run or stats["warnings"]:
                print(f"  {'[dry]' if args.dry_run else '[OK] '} {label:<60} "
                      f"consumo={stats['inserted_consumption']:4d} skip={stats['skipped_consumption']:3d} "
                      f"stock={stats['inserted_stock']:3d}  {status}")

    session.close()
    print(f"\n{'--- DRY RUN ---' if args.dry_run else '--- COMPLETADO ---'}")
    print(f"Registros de consumo: {total_consumed}")
    print(f"Registros de stock:   {total_stock}")
    print(f"Advertencias:         {total_warnings}")


if __name__ == "__main__":
    main()
